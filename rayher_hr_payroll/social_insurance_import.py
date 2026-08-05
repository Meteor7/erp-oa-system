from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from openpyxl import load_workbook

CENT = Decimal("0.01")
REQUIRED_HEADERS = {
    (2, 1): "序号",
    (2, 2): "姓名",
    (2, 3): "证件号码",
    (2, 4): "费款所属期起",
    (2, 5): "费款所属期止",
    (2, 9): "缴费工资",
}
AMOUNT_COLUMNS = {
    "receivable_total": 6,
    "employee_total": 7,
    "company_total": 8,
    "contribution_base": 9,
    "company_pension": 11,
    "employee_pension": 13,
    "company_medical": 15,
    "local_medical": 17,
    "employee_medical": 19,
    "company_maternity": 21,
    "company_injury": 27,
    "company_unemployment": 35,
    "employee_unemployment": 37,
    "local_pension": 39,
}


def _money(value):
    if value in (None, ""):
        return Decimal("0.00")
    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)


def canonical_period(value):
    text = str(value or "").strip().replace("/", "-")
    if len(text) == 7 and text[4] == "-":
        year, month = text.split("-")
        if year.isdigit() and month.isdigit() and 1 <= int(month) <= 12:
            return f"{int(year):04d}-{int(month):02d}"
    raise ValueError("社保期间格式错误")


def normalize_identity(value):
    return str(value or "").strip().replace(" ", "").upper()


def mask_identity(value):
    text = normalize_identity(value)
    if len(text) <= 8:
        return "*" * len(text)
    return f"{text[:4]}{'*' * (len(text) - 8)}{text[-4:]}"


def parse_social_workbook(content, period=None):
    expected_period = canonical_period(period) if period else None
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    if "申报明细" not in workbook.sheetnames:
        raise ValueError("缺少工作表：申报明细")
    sheet = workbook["申报明细"]
    for (row, column), expected in REQUIRED_HEADERS.items():
        if str(sheet.cell(row=row, column=column).value or "").strip() != expected:
            raise ValueError(f"社保Excel表头不匹配：缺少{expected}")

    active = False
    parsed = []
    detected_period = None
    for values in sheet.iter_rows(min_row=3, values_only=True):
        marker = str(values[0] or "").strip()
        label = str(values[1] or "").strip() if len(values) > 1 else ""
        if marker == "在职人员":
            active = True
            continue
        if active and (label == "小计" or marker in ("退休人员", "家属统筹人员")):
            break
        if not active or not marker:
            continue
        row_period_start = canonical_period(values[3])
        row_period_end = canonical_period(values[4])
        if row_period_start != row_period_end:
            raise ValueError(f"Excel明细起止期间不一致：{row_period_start}至{row_period_end}")
        if expected_period and row_period_start != expected_period:
            raise ValueError(f"Excel期间与所选期间不一致：{row_period_start}至{row_period_end}")
        if detected_period is None:
            detected_period = row_period_start
        elif row_period_start != detected_period:
            raise ValueError(f"Excel中存在多个社保期间：{detected_period}、{row_period_start}")
        row = {
            "row_number": int(marker),
            "excel_name": str(values[1] or "").strip(),
            "identity_number": normalize_identity(values[2]),
            "period_start": row_period_start,
            "period_end": row_period_end,
        }
        for key, column in AMOUNT_COLUMNS.items():
            row[key] = _money(values[column - 1] if len(values) >= column else None)
        parsed.append(row)
    if not parsed:
        raise ValueError("Excel中没有可读取的在职人员明细")
    return parsed


def normalize_name(value):
    return "".join(str(value or "").split())


def match_employee(row, employees):
    employees = list(employees or [])
    identity = normalize_identity(row.get("identity_number"))
    identity_matches = [
        employee for employee in employees
        if identity and normalize_identity(employee.get("id_card_number_cn")) == identity
    ]
    if len(identity_matches) == 1:
        return identity_matches[0], ""
    if len(identity_matches) > 1:
        return None, "证件号码关联到多名员工，无法自动判断"

    name = normalize_name(row.get("excel_name"))
    name_matches = [
        employee for employee in employees
        if name and normalize_name(employee.get("employee_name")) == name
    ]
    if len(name_matches) == 1:
        return name_matches[0], "证件号码未匹配，已按姓名唯一关联"
    if len(name_matches) > 1:
        return None, "证件号码未匹配且姓名不唯一，无法自动判断"
    return None, "证件号码和姓名均无法关联员工"

def reconcile_row(row, employee, company=None, association_reason=""):
    from rayher_hr_payroll.hr_contributions import (
        calculate_social_snapshot,
        compare_social_snapshot,
        household_classification,
        social_policy,
    )

    result = dict(row)
    reasons = [association_reason] if association_reason and "按姓名唯一关联" not in association_reason else []
    if not employee:
        result.update({
            "employee": None,
            "employee_name": None,
            "status": "Pending",
            "original_match_status": "Pending",
            "difference_reasons": association_reason or "证件号码无法关联员工",
        })
        return result

    employee_company = employee.get("company")
    employee_base = _money(employee.get("social_insurance_base"))
    base_source = str(employee.get("social_base_source") or "").strip()
    household_source = str(employee.get("household_registration") or "").strip()
    household = household_classification(household_source)
    policy = social_policy(employee_company, row.get("period_start"))

    if company and employee_company != company:
        reasons.append("关联员工所属公司与导入公司不一致")
    if str(employee.get("employee_name") or "").strip() != row["excel_name"]:
        reasons.append("证件号码一致但姓名不同")
    if employee_base <= 0:
        reasons.append("员工档案社保基数未维护")
    elif employee_base != _money(row.get("contribution_base")):
        reasons.append("Excel缴费工资与员工档案社保基数不一致")
    if household is None:
        reasons.append("员工户口信息未维护或无法判断深户/非深户")
    if policy is None:
        reasons.append("所属公司社保政策未配置")

    calculated = {}
    comparison = {"reasons": [], "differences": {"receivable_total": _money(0)}}
    if employee_base > 0 and household is not None and policy is not None:
        calculated = calculate_social_snapshot(employee_base, household, policy)
        comparison = compare_social_snapshot(row, calculated)
        reasons.extend(comparison["reasons"])

    result.update({
        "employee": employee.get("name"),
        "employee_name": employee.get("employee_name"),
        "company": employee_company,
        "reference_base": employee_base,
        "base_source": base_source,
        "household_source": household_source,
        "household_classification": "深户" if household is True else "非深户" if household is False else "待完善",
        "policy_version": (policy or {}).get("version", ""),
        "calculated": calculated,
        "differences": comparison["differences"],
        "difference_total": comparison["differences"].get("receivable_total", _money(0)),
    })
    matched = not reasons
    result["status"] = "Matched" if matched else "Pending"
    result["original_match_status"] = result["status"]
    result["difference_reasons"] = "；".join(dict.fromkeys(reasons))
    return result


def select_workbench_batches(batches, requested_batch=None):
    batches = list(batches or [])
    if requested_batch:
        return [row["name"] for row in batches if row.get("name") == requested_batch]
    preview = [row["name"] for row in batches if row.get("status") == "Pending Activation"]
    if preview:
        return preview[:1]
    return [row["name"] for row in batches if row.get("is_current")][:1]

def can_activate_rows(rows):
    rows = list(rows or [])
    return bool(rows) and all(
        row.get("employee") and row.get("status") in ("Matched", "Confirmed Difference")
        for row in rows
    )

def summarize_rows(rows):
    rows = list(rows or [])
    total = len(rows)
    linked = sum(1 for row in rows if row.get("employee"))
    comparable = linked
    matched = sum(1 for row in rows if row.get("original_match_status") == "Matched")
    pending = sum(1 for row in rows if row.get("status") == "Pending")
    confirmed = sum(1 for row in rows if row.get("status") == "Confirmed Difference")

    def rate(numerator, denominator):
        return round(numerator * 100 / denominator, 2) if denominator else 0.0

    def total_money(field):
        return sum((_money(row.get(field)) for row in rows), Decimal("0.00"))

    return {
        "total_rows": total,
        "linked_rows": linked,
        "comparable_rows": comparable,
        "matched_rows": matched,
        "pending_rows": pending,
        "confirmed_difference_rows": confirmed,
        "unlinked_rows": total - linked,
        "match_rate": rate(matched, comparable),
        "comparable_rate": rate(comparable, total),
        "automatic_link_rate": rate(linked, total),
        "employee_total": total_money("employee_total"),
        "company_total": total_money("company_total"),
        "receivable_total": total_money("receivable_total"),
    }
