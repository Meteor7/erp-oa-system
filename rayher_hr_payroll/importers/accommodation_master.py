from datetime import date, datetime, timedelta
from pathlib import Path

import frappe
from rayher_hr_payroll import site_time

SPECIAL_PURPOSES = {"空", "客房", "午休室"}
SUMMARY_LABELS = {"合计", "小计"}


def _text(value):
    return str(value or "").strip()


def _amount(value):
    return float(value or 0)


def _date_value(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def parse_master_row(row):
    values = list(row[:12]) + [None] * max(0, 12 - len(row))
    (
        room_no,
        room_type,
        lease_start_date,
        lease_end_date,
        monthly_rent,
        paid_deposit,
        rent_contract_attachment,
        deposit_attachment,
        occupant_label,
        start_date,
        end_date,
        remark,
    ) = values[:12]
    room_no = _text(room_no)
    occupant_label = _text(occupant_label)
    room_purpose = occupant_label if occupant_label in SPECIAL_PURPOSES else ""
    room = {
        "room_no": room_no,
        "room_type": _text(room_type),
        "lease_start_date": _date_value(lease_start_date),
        "lease_end_date": _date_value(lease_end_date),
        "monthly_rent": _amount(monthly_rent),
        "default_rent": _amount(monthly_rent),
        "paid_deposit": _amount(paid_deposit),
        "rent_contract_attachment": _text(rent_contract_attachment),
        "deposit_attachment": _text(deposit_attachment),
        "occupant_label": occupant_label,
        "room_purpose": room_purpose,
        "remark": _text(remark),
        "is_active": 1,
    }
    occupancy = None
    if occupant_label and not room_purpose:
        occupancy = {
            "employee_name": occupant_label,
            "start_date": _date_value(start_date),
            "end_date": _date_value(end_date),
            "remark": _text(remark),
        }
    return {"room": room, "occupancy": occupancy}


def occupancy_is_effective(start_date, end_date, period_start, period_end):
    start_date = _date_value(start_date)
    end_date = _date_value(end_date)
    period_start = _date_value(period_start)
    period_end = _date_value(period_end)
    return (not start_date or start_date <= period_end) and (not end_date or end_date >= period_start)


def _employee_for_name(employee_name):
    # 只匹配在职员工：离职员工的同名新人不能被登记到离职档上（修复 2026-09-29，模块 7）。
    matches = frappe.get_all(
        "Employee",
        filters={"employee_name": employee_name, "status": "Active"},
        pluck="name",
        limit=2,
    )
    return matches[0] if len(matches) == 1 else None


def _upsert_room(values):
    name = frappe.db.get_value("Accommodation Room", {"room_no": values["room_no"]}, "name")
    doc = frappe.get_doc("Accommodation Room", name) if name else frappe.new_doc("Accommodation Room")
    doc.update(values)
    doc.save(ignore_permissions=True)
    return bool(name)


def _move_out_date(old_start, new_start):
    move_out = new_start - timedelta(days=1) if new_start else site_time.today()
    old_start = _date_value(old_start)
    return max(move_out, old_start) if old_start else move_out


def _upsert_occupancy(room_no, values, employee):
    existing = frappe.db.get_value(
        "Accommodation Occupancy",
        {"employee": employee, "room": room_no},
        "name",
    )
    if not existing:
        # 修复（2026-09-29，审查 #15a）：换房时原来只把旧记录 is_active 改成 0，唯一键
        # active_employee_key 还在，新记录保存撞唯一键，整个导入失败；旧记录也没有退宿日期，
        # 下次保存又变回"在住"。现在按正常退宿保存旧记录：退宿日 = 新入住日前一天（无入住日按今天）。
        old_current = frappe.get_all(
            "Accommodation Occupancy",
            filters={"employee": employee, "is_active": 1},
            pluck="name",
        )
        for name in old_current:
            old = frappe.get_doc("Accommodation Occupancy", name)
            old.end_date = _move_out_date(old.start_date, values["start_date"])
            old.save(ignore_permissions=True)
    doc = frappe.get_doc("Accommodation Occupancy", existing) if existing else frappe.new_doc("Accommodation Occupancy")
    doc.update(
        {
            "employee": employee,
            "room": room_no,
            "start_date": values["start_date"],
            "end_date": values["end_date"],
            "remark": values["remark"],
        }
    )
    doc.save(ignore_permissions=True)
    return bool(existing)


def import_master(workbook_path):
    """Import real room and occupancy master data without creating monthly fees."""
    from openpyxl import load_workbook

    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    sheet = load_workbook(path, data_only=True, read_only=True)["宿舍基础资料"]
    result = {
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "occupancy_created": 0,
        "occupancy_updated": 0,
        "unmatched": [],
        "errors": [],
        "monthly_rent_total": 0.0,
        "paid_deposit_total": 0.0,
    }
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        room_no = _text(raw_row[0] if raw_row else None)
        if not room_no or room_no in SUMMARY_LABELS:
            result["skipped"] += 1
            continue
        parsed = parse_master_row(raw_row)
        existed = _upsert_room(parsed["room"])
        result["updated" if existed else "created"] += 1
        result["monthly_rent_total"] += parsed["room"]["monthly_rent"]
        result["paid_deposit_total"] += parsed["room"]["paid_deposit"]
        occupancy = parsed["occupancy"]
        if occupancy:
            employee = _employee_for_name(occupancy["employee_name"])
            if not employee:
                result["unmatched"].append(occupancy["employee_name"])
            else:
                # 一行出错（例如房间已有别人在住）只跳过这一行，不影响其他房间。
                frappe.db.savepoint("accommodation_row")
                try:
                    existed = _upsert_occupancy(room_no, occupancy, employee)
                except Exception as exc:  # noqa: BLE001 - reported per row
                    frappe.db.rollback(save_point="accommodation_row")
                    result["errors"].append(f"{room_no} {occupancy['employee_name']}: {exc}")
                else:
                    result["occupancy_updated" if existed else "occupancy_created"] += 1
    frappe.db.commit()
    return result


def import_rooms(workbook_path):
    """Backward-compatible entry point for the complete master import."""
    return import_master(workbook_path)


def audit_counts():
    doctypes = (
        "Accommodation Room",
        "Accommodation Occupancy",
        "Accommodation Deduction",
        "Employee Social Insurance",
        "Employee Housing Fund",
        "Salary Slip",
    )
    return {doctype: frappe.db.count(doctype) for doctype in doctypes}
